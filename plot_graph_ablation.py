import shutil
import warnings
import zipfile
from xml.etree import ElementTree as ET

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import seaborn as sns
from matplotlib.patches import Patch

try:
    import openpyxl
except ImportError:
    openpyxl = None

warnings.filterwarnings(
    "ignore",
    category=FutureWarning,
    message=".*use_inf_as_na option is deprecated.*",
)

# ==========================
# CONFIG
# ==========================
XLSX_PATH = "graph_results.xlsx"
SHEET_NAME = "baseline_comparison_0606"
OUT_PDF = "plot_baseline_comparison0606.pdf"
OUT_PNG = "plot_baseline_comparison0606.png"

N_COLS = 3
N_ROWS = 2
ROW_HEIGHT = 1.95
AXIS_COLOR = "black"
TICK_LENGTH = 3
TICK_WIDTH = 0.6

DATASET_ORDER = ["Asia", "Sachs", "Alarm", "Insurance", "Hailfinder", "SIIM"]
DATASET_ALIASES = {"Siim": "SIIM"}

METHOD_SPECS = [
    ("Loc.", "Loc."),
    ("S-Union", "S-Union"),
    ("S-Intersection", "S-Inters."),
    ("S-BN fusion", "S-BN fusion"),
    ("S-F-CMs", "S-F-CMs"),
    ("Union", "Union"),
    ("Intersection", "Inters."),
    ("BN fusion", "BN-fusion"),
    ("F-CMs (Ours)", "F-CMs (Ours)"),
]
METHOD_ORDER = [m[0] for m in METHOD_SPECS]
METHOD_X_LABELS = [m[1] for m in METHOD_SPECS]
METHOD_IGNORE = {"S-Majority", "Majority"}


def setup_icml_plot(two_column=True):
    """ICML/NeurIPS-like typography + seaborn styling."""
    figure_width = 7 if two_column else 3.5
    use_tex = shutil.which("latex") is not None
    if not use_tex:
        print("[warn] LaTeX not found; falling back to Matplotlib text rendering.")
    rc = {
        "text.usetex": use_tex,
        "font.family": "serif",
        "font.serif": ["Times New Roman", "Times", "Computer Modern Roman"],
        "axes.labelsize": 9,
        "axes.titlesize": 9,
        "legend.fontsize": 8,
        "xtick.labelsize": 8,
        "ytick.labelsize": 8,
        "lines.linewidth": 1.2,
        "lines.markersize": 3,
        "figure.figsize": (figure_width, figure_width * 0.18),
        "figure.dpi": 300,
        "axes.grid": True,
        "grid.alpha": 0.3,
        "grid.linestyle": "--",
        "legend.frameon": False,
        "pdf.fonttype": 42,
        "ps.fonttype": 42,
        "axes.edgecolor": AXIS_COLOR,
        "axes.labelcolor": AXIS_COLOR,
        "xtick.color": AXIS_COLOR,
        "ytick.color": AXIS_COLOR,
        "text.color": AXIS_COLOR,
    }
    sns.set_theme(context="paper", style="whitegrid", palette="deep", rc=rc)
    return figure_width, use_tex


PAPER_WIDTH, USE_TEX = setup_icml_plot(two_column=True)
PALETTE = sns.color_palette("deep", n_colors=3)
COLOR_LOCAL = PALETTE[0]
COLOR_STATIC = PALETTE[2]
COLOR_DYNAMIC = PALETTE[1]
OURS_X_LABELS = {"S-F-CMs", "F-CMs", "F-CMs (Ours)"}


def _format_xtick_label(label):
    if USE_TEX and label in OURS_X_LABELS:
        return rf"\textbf{{{label}}}"
    return label


METHOD_X_LABELS_PLOT = [_format_xtick_label(lbl) for lbl in METHOD_X_LABELS]


def _normalize_dataset_name(name):
    if name is None:
        return None
    raw = str(name).strip()
    return DATASET_ALIASES.get(raw, raw)


def _to_float(value):
    if value is None:
        return np.nan
    try:
        return float(value)
    except (TypeError, ValueError):
        return np.nan


def _column_to_index(col_ref):
    idx = 0
    for ch in col_ref:
        if not ch.isalpha():
            break
        idx = idx * 26 + (ord(ch.upper()) - ord("A") + 1)
    return idx


def _load_shared_strings(zf):
    if "xl/sharedStrings.xml" not in zf.namelist():
        return []
    root = ET.fromstring(zf.read("xl/sharedStrings.xml"))
    ns = {"a": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    shared = []
    for si in root.findall("a:si", ns):
        chunks = []
        for t in si.iter("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}t"):
            chunks.append(t.text or "")
        shared.append("".join(chunks))
    return shared


def _find_sheet_xml_path(zf, sheet_name):
    ns_wb = {
        "a": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
        "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    }
    wb = ET.fromstring(zf.read("xl/workbook.xml"))
    rid = None
    for sheet in wb.find("a:sheets", ns_wb):
        if sheet.attrib.get("name") == sheet_name:
            rid = sheet.attrib.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
            break
    if rid is None:
        raise RuntimeError(f"Sheet '{sheet_name}' not found in {XLSX_PATH}.")

    rels = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
    for rel in rels:
        if rel.attrib.get("Id") == rid:
            target = rel.attrib.get("Target")
            if target.startswith("/"):
                target = target[1:]
            return f"xl/{target}" if not target.startswith("xl/") else target
    raise RuntimeError(f"Could not resolve XML path for sheet '{sheet_name}'.")


def _read_sheet_matrix_from_xml(xlsx_path, sheet_name):
    ns = {"a": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    matrix = {}
    max_row = 0
    max_col = 0

    with zipfile.ZipFile(xlsx_path) as zf:
        shared = _load_shared_strings(zf)
        sheet_xml = _find_sheet_xml_path(zf, sheet_name)
        root = ET.fromstring(zf.read(sheet_xml))
        sheet_data = root.find("a:sheetData", ns)
        if sheet_data is None:
            return matrix, max_row, max_col

        for row in sheet_data.findall("a:row", ns):
            r = int(row.attrib["r"])
            max_row = max(max_row, r)
            for cell in row.findall("a:c", ns):
                ref = cell.attrib.get("r", "")
                col_ref = "".join(ch for ch in ref if ch.isalpha())
                c = _column_to_index(col_ref)
                if c == 0:
                    continue
                max_col = max(max_col, c)

                value = None
                c_type = cell.attrib.get("t")
                if c_type == "inlineStr":
                    is_node = cell.find("a:is", ns)
                    if is_node is not None:
                        text_nodes = is_node.findall(".//a:t", ns)
                        value = "".join((t.text or "") for t in text_nodes)
                else:
                    v_node = cell.find("a:v", ns)
                    if v_node is not None:
                        raw = v_node.text
                        if c_type == "s":
                            idx = int(raw)
                            value = shared[idx] if 0 <= idx < len(shared) else None
                        else:
                            value = raw

                if value is not None and value != "":
                    matrix[(r, c)] = value

    return matrix, max_row, max_col


def _parse_baseline_rows(matrix, max_row, max_col):
    header_pos = None
    for r in range(1, max_row + 1):
        for c in range(1, max_col + 1):
            if str(matrix.get((r, c), "")).strip() == "Datasets":
                header_pos = (r, c)
                break
        if header_pos is not None:
            break

    if header_pos is None:
        raise RuntimeError("Could not locate 'Datasets' header in baseline_comparison sheet.")

    _, dataset_col = header_pos
    current_dataset = None
    records = []

    for r in range(header_pos[0], max_row + 1):
        dataset_val = matrix.get((r, dataset_col + 1))
        method_val = matrix.get((r, dataset_col + 2))
        mean_val = matrix.get((r, dataset_col + 3))
        std_val = matrix.get((r, dataset_col + 4))

        if dataset_val is not None:
            current_dataset = _normalize_dataset_name(dataset_val)

        if method_val is None or current_dataset is None:
            continue

        method = str(method_val).strip()
        if method in METHOD_IGNORE:
            continue

        mean = _to_float(mean_val)
        std = _to_float(std_val)
        if np.isnan(mean):
            continue

        records.append(
            {
                "dataset": current_dataset,
                "method": method,
                "mean": mean,
                "std": 0.0 if np.isnan(std) else std,
            }
        )

    df = pd.DataFrame(records)
    if df.empty:
        raise RuntimeError("No baseline records found in baseline_comparison sheet.")
    return df


def parse_baseline_comparison(xlsx_path, sheet_name):
    if openpyxl is not None:
        wb = openpyxl.load_workbook(xlsx_path, data_only=True)
        if sheet_name not in wb.sheetnames:
            raise RuntimeError(f"Sheet '{sheet_name}' not found in {xlsx_path}.")
        ws = wb[sheet_name]

        header_pos = None
        for r in range(1, ws.max_row + 1):
            for c in range(1, ws.max_column + 1):
                val = ws.cell(r, c).value
                if isinstance(val, str) and val.strip() == "Datasets":
                    header_pos = (r, c)
                    break
            if header_pos is not None:
                break

        if header_pos is None:
            raise RuntimeError("Could not locate 'Datasets' header in baseline_comparison sheet.")

        _, dataset_col = header_pos
        current_dataset = None
        records = []

        for r in range(header_pos[0], ws.max_row + 1):
            dataset_val = ws.cell(r, dataset_col + 1).value
            method_val = ws.cell(r, dataset_col + 2).value
            mean_val = ws.cell(r, dataset_col + 3).value
            std_val = ws.cell(r, dataset_col + 4).value

            if dataset_val is not None:
                current_dataset = _normalize_dataset_name(dataset_val)
            if method_val is None or current_dataset is None:
                continue

            method = str(method_val).strip()
            if method in METHOD_IGNORE:
                continue

            mean = _to_float(mean_val)
            std = _to_float(std_val)
            if np.isnan(mean):
                continue

            records.append(
                {
                    "dataset": current_dataset,
                    "method": method,
                    "mean": mean,
                    "std": 0.0 if np.isnan(std) else std,
                }
            )

        df = pd.DataFrame(records)
        if df.empty:
            raise RuntimeError("No baseline records found in baseline_comparison sheet.")
        return df

    matrix, max_row, max_col = _read_sheet_matrix_from_xml(xlsx_path, sheet_name)
    return _parse_baseline_rows(matrix, max_row, max_col)


def method_group(method_name):
    if method_name == "Loc.":
        return "local"
    if method_name.startswith("S-"):
        return "static"
    return "dynamic"


def method_color(method_name):
    group = method_group(method_name)
    if group == "local":
        return COLOR_LOCAL
    if group == "static":
        return COLOR_STATIC
    return COLOR_DYNAMIC


def dataset_sort_key(name):
    if name in DATASET_ORDER:
        return (0, DATASET_ORDER.index(name))
    return (1, str(name))


def build_figure(df):
    datasets = sorted(df["dataset"].unique(), key=dataset_sort_key)

    if len(datasets) != 6:
        print(f"[warn] Expected 6 datasets, found {len(datasets)}: {datasets}")

    fig_w = PAPER_WIDTH
    fig_h = max(2.6, ROW_HEIGHT * N_ROWS)
    fig, axes = plt.subplots(
        N_ROWS,
        N_COLS,
        figsize=(fig_w, fig_h),
        squeeze=False,
    )
    axes_flat = axes.flatten()

    x = np.arange(len(METHOD_ORDER))

    for idx, dataset in enumerate(datasets):
        if idx >= len(axes_flat):
            break
        ax = axes_flat[idx]
        sub = df[df["dataset"] == dataset]

        means = []
        stds = []
        for method in METHOD_ORDER:
            row = sub[sub["method"] == method]
            if row.empty:
                means.append(np.nan)
                stds.append(np.nan)
            else:
                means.append(float(row.iloc[0]["mean"]))
                stds.append(float(row.iloc[0]["std"]))

        means = np.array(means, dtype=float)
        stds = np.array(stds, dtype=float)
        bar_colors = [method_color(m) for m in METHOD_ORDER]

        bars = ax.bar(
            x,
            np.nan_to_num(means, nan=0.0),
            yerr=np.nan_to_num(stds, nan=0.0),
            color=bar_colors,
            edgecolor=AXIS_COLOR,
            linewidth=0.5,
            error_kw={"elinewidth": 0.8, "capsize": 2.5, "capthick": 0.8},
        )

        for j, m in enumerate(means):
            if np.isnan(m):
                bars[j].set_alpha(0.2)
                bars[j].set_hatch("//")

        ax.set_title(dataset, pad=4)
        ax.set_xticks(x)
        ticklabels = ax.set_xticklabels(METHOD_X_LABELS_PLOT, rotation=30, ha="right")
        if not USE_TEX:
            for tick in ticklabels:
                if tick.get_text() in OURS_X_LABELS:
                    tick.set_fontweight("bold")

        col_idx = idx % N_COLS
        if col_idx == 0:
            ax.set_ylabel("Diff. pairs (↓)")
        else:
            ax.set_ylabel("")

        finite_means = means[np.isfinite(means)]
        finite_stds = stds[np.isfinite(stds)]
        if finite_means.size > 0:
            y_top = np.max(finite_means + finite_stds)
            pad = max(1.0, 0.08 * y_top)
            ax.set_ylim(0, y_top + pad)

        ax.grid(True, axis="y", linewidth=0.4, alpha=0.3, linestyle="--")
        ax.grid(False, axis="x")
        ax.tick_params(
            axis="both",
            which="major",
            length=TICK_LENGTH,
            width=TICK_WIDTH,
            direction="out",
            colors=AXIS_COLOR,
            bottom=True,
            left=True,
        )
        for spine in ("left", "bottom"):
            ax.spines[spine].set_color(AXIS_COLOR)

        # Keep right/top spines hidden, consistent with seaborn paper style.
        ax.spines["right"].set_visible(False)
        ax.spines["top"].set_visible(False)

    # Hide unused subplot slots if needed.
    for idx in range(len(datasets), len(axes_flat)):
        axes_flat[idx].set_visible(False)

    legend_handles = [
        Patch(facecolor=COLOR_LOCAL, edgecolor=AXIS_COLOR, label="Local (Loc.)"),
        Patch(facecolor=COLOR_STATIC, edgecolor=AXIS_COLOR, label="Static view (S-*)"),
        Patch(facecolor=COLOR_DYNAMIC, edgecolor=AXIS_COLOR, label="Dynamic view"),
    ]
    fig.legend(
        handles=legend_handles,
        loc="upper center",
        bbox_to_anchor=(0.5, 1.01),
        ncol=3,
        frameon=False,
    )

    fig.tight_layout(rect=[0, 0, 1, 0.95], h_pad=1.2, w_pad=0.8)
    return fig


def main():
    df = parse_baseline_comparison(XLSX_PATH, SHEET_NAME)

    # Keep only requested methods and preserve order.
    df = df[df["method"].isin(METHOD_ORDER)].copy()
    df["method"] = pd.Categorical(df["method"], categories=METHOD_ORDER, ordered=True)
    df = df.sort_values(["dataset", "method"])

    fig = build_figure(df)
    fig.savefig(OUT_PDF, bbox_inches="tight")
    fig.savefig(OUT_PNG, dpi=300, bbox_inches="tight")
    print(f"[OK] Saved: {OUT_PDF}")
    print(f"[OK] Saved: {OUT_PNG}")


if __name__ == "__main__":
    main()
